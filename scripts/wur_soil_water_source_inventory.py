"""M1-35: audit original WUR groundwater and soil-water-potential XLSX safely.

ten Den et al. (2024), Harvard Dataverse v3.0 CC BY-NC-SA 4.0.
Outputs ONLY sheet-level structural/typed counts and controlled clue categories.
Do not export original headers, values, daily series, sample identifiers or plot IDs.
Schema presence != retention curve, soil hydraulic parameters or WOFOST validation.
"""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import date,datetime
import hashlib,io,json,math,re
from pathlib import Path

DOI="10.7910/DVN/1LC6W7"
LICENSE="CC BY-NC-SA 4.0"
ORIGINALS={
    "groundwater":{"file_id":10284788,"filename":"Groundwaterlevel.xlsx","sha256":"7c546ba5493fa28bb073073071ff44be562e89d4401467caefe73b82911f65de"},
    "waterpotential":{"file_id":10284794,"filename":"Waterpotential.xlsx","sha256":"7a6924d01eff020215a1aca9c7de36fa4d97b33a0e99cfcf330127bedcfdec15"},
}
CLUES={
  "date_time":r"\bdate\b|\bdatum\b|\btime\b|\btimestamp\b|\bday\b",
  "groundwater":r"\bgroundwater\b|\bground\s*water\b|\bwater\s*table\b|\bgrondwater\b",
  "potential":r"\bpotential\b|\bwaterpotentiaal\b|\bzuigspanning\b|\bsuction\b",
  "soil_moisture":r"\bsoil\s*moisture\b|\bvolumetric\b|\btheta\b|\bwater\s*content\b",
  "depth":r"\bdepth\b|\bdiepte\b|\bbelow\b|\bfield\s*level\b",
  "depth_cm":r"(?<![a-z])cm(?![a-z])|\bcentimet(?:er|re)s?\b",
  "depth_m":r"(?<![a-z])m(?![a-z])|\bmeters?\b|\bmetres?\b",
  "tension_pF":r"(?<![a-z])pf(?![a-z])",
  "pressure_kPa":r"(?<![a-z])kpa(?![a-z])",
  "pressure_hPa":r"(?<![a-z])hpa(?![a-z])",
  "pressure_cbar":r"(?<![a-z])cbar(?![a-z])|\bcentibar",
  "pressure_mbar":r"(?<![a-z])mbar(?![a-z])|\bmillibar",
  "site":r"\blocation\b|\bsite\b|\blocatie\b",
  "year":r"\byear\b|\bjaar\b",
  "plot":r"\bplot\b|\bplotnumber\b|\bperceel\b",
  "treatment":r"\btreatment\b|\bw1\b|\bw2\b|\birrigation\b",
}
MAX_ROWS=200000
MAX_COLS=140
def classify(v):
    if v is None or isinstance(v,str) and v.strip().upper() in ("","NA","NAN","NULL"):
        return "blank_or_missing"
    if isinstance(v,(datetime,date)):return "typed_date"
    if isinstance(v,bool):return "boolean"
    if isinstance(v,(int,float)):
        return "numeric" if math.isfinite(v) else "invalid_numeric"
    if isinstance(v,str):return "text"
    return "unrecognized"

def checked_date(v):
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    # No locale guesses: ISO only.
    if isinstance(v,str):
        try:
            d=date.fromisoformat(v.strip())
            return d if v.strip()==d.isoformat() else None
        except ValueError:return None
    return None

def metadata_tags(v):
    if not isinstance(v,str):return []
    return sorted(k for k,pattern in CLUES.items() if re.search(pattern,v,re.IGNORECASE))

def inspect_sheet(ws,*,max_rows=MAX_ROWS):
    type_counts=Counter()
    year_dates=Counter()
    tagged={}
    used=0
    columns=ws.max_column
    if columns is None or columns<1 or columns>MAX_COLS:
        raise ValueError("Original soil workbook columns exceed safety bounds")
    for row_number,row in enumerate(ws.iter_rows(values_only=True),start=1):
        if row_number>max_rows:raise ValueError("Soil source rows exceed safety limit")
        used+=1
        for column_number,v in enumerate(row,start=1):
            if column_number>MAX_COLS:raise ValueError("Unexpected very wide source row")
            type_counts[classify(v)]+=1
            d=checked_date(v)
            if d and d.year in (2018,2019,2020,2021):
                year_dates[str(d.year)]+=1
            if row_number<=35:
                labels=metadata_tags(v)
                if labels:
                    tagged.setdefault(str(row_number),[]).append({
                        "source_column_one_based":column_number,
                        "controlled_clues":labels,
                    })
    if used==0:raise ValueError("Original soil source sheet empty")
    return {
        "physical_rows":used,
        "physical_columns":columns,
        "cell_type_counts":dict(type_counts),
        "typed_or_ISO_date_cells_by_year":dict(year_dates),
        "controlled_metadata_clues_by_position_first_35_rows":tagged,
        "original_label_text_or_values_exported":0,
        "site_year_mapping_independently_verified":False,
        "units_per_physical_column_approved":False,
        "water_retention_curve_derived":False,
        "saturated_hydraulic_conductivity_verified":False,
        "soil_depth_horizon_mapping_approved":False,
    }

def inspect_workbook(blob,kind,*,allow_bootstrap=False):
    spec=ORIGINALS.get(kind)
    if spec is None:raise ValueError("Unknown WUR original kind")
    digest=hashlib.sha256(blob).hexdigest()
    pinned=spec["sha256"]
    if pinned is None and not allow_bootstrap:
        raise ValueError("Original WUR soil source needs independently pinned SHA256")
    if pinned is not None and digest!=pinned:
        raise ValueError("Original WUR soil workbook SHA256 mismatch")
    from openpyxl import load_workbook
    try:book=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    except Exception as exc:raise ValueError("Original soil workbook not valid XLSX") from exc
    try:
        details=[]
        for ordinal,ws in enumerate(book.worksheets,start=1):
            details.append({
                "source_sheet_one_based":ordinal,
                "source_sheet_type":"meta" if ws.title.strip().lower()=="meta" else "data_or_other",
                **inspect_sheet(ws),
            })
    finally:book.close()
    if not details:raise ValueError("No original soil worksheets")
    return {
        "kind":kind,
        "source_file_id":spec["file_id"],
        "source_filename":spec["filename"],
        "source_sha256":digest,
        "source_digest_independently_pinned":pinned is not None,
        "number_of_sheets":len(details),
        "sheet_structures":details,
        "individual_source_observations_exported":0,
        "native_file_republished":False,
        "field_soil_parameters_derived":False,
    }

def assess_sources(data,*,allow_bootstrap=False):
    if set(data)!=set(ORIGINALS):
        raise ValueError("Both WUR original soil observation files required")
    results=[inspect_workbook(data[k],k,allow_bootstrap=allow_bootstrap)
             for k in ORIGINALS]
    return {
        "status":"M1_35_WUR_WATER_OBSERVATIONS_SOURCE_INVENTORY_NOT_SOIL_PARAMETERIZATION",
        "source_doi":DOI,
        "source_version":"3.0",
        "source_license":LICENSE,
        "sources":results,
        "water_table_observations_are_not_hydraulic_soil_curves":True,
        "soil_potential_observations_are_not_crop_soil_hydraulic_parameters":True,
        "water_units_sensor_depth_and_site_join_review_required":True,
        "2019_development_2020_biomass_holdout_not_used":True,
        "WOFOST_field_soil_provider_approved":False,
        "WOFOST_field_crop_run_approved":False,
        "field_yield_errors_computed":False,
        "raw_individual_soil_values_exported":0,
        "raw_header_text_exported":0,
        "limits":[
          "Controlled vocabulary clues are only labels, not physical unit certification.",
          "Water potential and groundwater level are observed environmental states; hydraulic soil response needs independent laboratory/field parameterization.",
          "Original soil files and sample records stay on a temporary CI runner.",
          "No irrigation blank may be silently imputed as zero; M1-34 evidence remains incomplete.",
        ],
    }

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--groundwater",required=True,type=Path)
    p.add_argument("--waterpotential",required=True,type=Path)
    p.add_argument("--out",required=True,type=Path)
    p.add_argument("--bootstrap-unpinned-inventory",action="store_true")
    args=p.parse_args()
    result=assess_sources({
        "groundwater":args.groundwater.read_bytes(),
        "waterpotential":args.waterpotential.read_bytes(),
    },allow_bootstrap=args.bootstrap_unpinned_inventory)
    args.out.mkdir(parents=True,exist_ok=True)
    (args.out/"m1_35_wur_soil_observations_source_inventory.json").write_text(
        json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps({
        "status":result["status"],
        "files":[{"kind":x["kind"],"sha256":x["source_sha256"],
                  "sheets":x["number_of_sheets"],
                  "summaries":[{"rows":y["physical_rows"],"columns":y["physical_columns"],
                                "year_date_counts":y["typed_or_ISO_date_cells_by_year"],
                                "controlled_clues":y["controlled_metadata_clues_by_position_first_35_rows"]}
                               for y in x["sheet_structures"]]}
                 for x in result["sources"]],
        "field_soil_ready":False,
    },sort_keys=True))
if __name__=="__main__":main()
