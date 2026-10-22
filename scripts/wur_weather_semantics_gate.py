"""M1-24: source-verified WUR weather missingness/metadata, NOT model validation.

ten Den et al. 2024, doi:10.7910/DVN/1LC6W7, v3.0, CC BY-NC-SA 4.0.
Exports positional counts and tentative controlled-vocabulary tags only, NOT
original worksheets, source header text, dates or weather measurements.
"""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import date
import hashlib
import io
import json
import math
from pathlib import Path
import re
from scripts.wur_weather_daily_gate import parse_day

EXPECTED_SHA = {
    "Lelystad": "ad66a23468a969b5795a656fb6989499e790a3e18179b9e3158f203b12ab8636",
    "Vredepeel": "0bb40a5aa91ae6bfd02a8e8b4f825e4d834207884c1f2f44caa12fd0e88a1fd6",
}
ORIGINAL_YEARS={2018:365,2019:365,2020:366}
NA={"","NA","NULL","NAN","N/A","-"}
# Only clues, never an approved field mapping.
FIELD_CLUES={
    "radiation":r"radiat|straling|irradiat",
    "rain":r"precip|rain|neerslag|regen",
    "temperature":r"temp|temperatuur|tmin|tmax",
    "humidity":r"humid|vochtigheid|\brh\b",
    "vapor_pressure":r"vapou?r pressure|dampdruk",
    "wind":r"\bwind",
    "evapotranspiration":r"evap|makkink|verdamp",
}
UNIT_CLUES={
    "mm":r"\bmm\b","cm":r"\bcm\b",
    "MJ_m2":r"\bmj\s*/?\s*m(?:2|²)\b",
    "kJ_m2":r"\bkj\s*/?\s*m(?:2|²)\b",
    "hPa":r"\bhpa\b","kPa":r"\bkpa\b",
    "percent":r"%","degC":r"(?:°\s*c|\bcelsius\b)",
    "m_per_s":r"\bm\s*/\s*s\b",
    "wind_at_2m":r"\b2\s*m(?:eter|etre)?\b",
}

def classify(value):
    if value is None or isinstance(value,str) and value.strip().upper() in NA:
        return "missing"
    if isinstance(value,bool):
        return "untyped"
    if isinstance(value,(int,float)):
        return "numeric" if math.isfinite(value) else "untyped"
    return "untyped"

def maximum_non_numeric_streak(values):
    longest=cur=0
    for v in values:
        cur=cur+1 if classify(v)!="numeric" else 0
        longest=max(longest,cur)
    return longest

def audit_sheet(ws,expected_years=None):
    rows=list(ws.iter_rows(min_row=1,values_only=True))
    if not rows:
        raise ValueError("Empty weather worksheet")
    first=None
    dated=[]
    for i,row in enumerate(rows):
        day=parse_day(row[0]) if row else None
        if day is None or day.year < 1950 or day.year > 2100:
            if first is not None:
                raise ValueError("Non-date inside source chronological series")
            continue
        if first is None:first=i
        dated.append((day,row))
    if not dated or first is None or not 2<=first<=20:
        raise ValueError("Missing original weather dates or metadata prefix")
    days=[d for d,_ in dated]
    if len(days)!=len(set(days)):
        raise ValueError("Duplicate original day")
    if days!=sorted(days) or (days[-1]-days[0]).days+1!=len(days):
        raise ValueError("Gap or out-of-order original day")
    totals=Counter(x.year for x in days)
    if expected_years is not None and dict(totals)!=dict(expected_years):
        raise ValueError("Original date window changed")
    ncol=max(len(r) for r in rows)
    prefix=rows[:first]
    cols=[]
    for i in range(1,ncol):
        meta=" ".join(row[i].strip().lower() for row in prefix
                      if len(row)>i and isinstance(row[i],str))
        by_year={}
        by_quarter={}
        for y in (2019,2020):
            values=[row[i] if len(row)>i else None for day,row in dated if day.year==y]
            kinds=Counter(classify(v) for v in values)
            by_year[str(y)]={
                "days":len(values),"finite_numeric":kinds["numeric"],
                "missing_na":kinds["missing"],"untyped_nonfinite":kinds["untyped"],
                "max_non_numeric_streak":maximum_non_numeric_streak(values)}
            for quarter in range(1,5):
                vals=[row[i] if len(row)>i else None for day,row in dated
                      if day.year==y and (day.month-1)//3+1==quarter]
                kind=Counter(classify(v) for v in vals)
                by_quarter[f"{y}Q{quarter}"]={
                    "days":len(vals),"finite_numeric":kind["numeric"],
                    "missing_na":kind["missing"],"untyped_nonfinite":kind["untyped"]}
        cols.append({
            "source_column_number":i+1,
            "original_header_text_exported":False,
            "candidate_variable_words_NOT_verified":[k for k,p in FIELD_CLUES.items()
                                                     if re.search(p,meta,re.I)],
            "candidate_unit_words_NOT_verified":[k for k,p in UNIT_CLUES.items()
                                                 if re.search(p,meta,re.I)],
            "counts_by_year":by_year,
            "counts_by_quarter":by_quarter,
            "approved_pcse_field":None,
            "instrument_method_and_unit_verified":False,
        })
    return {
        "original_daily_dates":len(days),
        "daily_dates_by_year":{str(y):n for y,n in sorted(totals.items())},
        "metadata_prefix_row_count":first,
        "date_sequence_complete":True,
        "columns":cols,"pcse_ready":False,
    }

def audit_originals(blobs):
    if set(blobs)!=set(EXPECTED_SHA):
        raise ValueError("Both WUR site workbooks required")
    for name,sha in EXPECTED_SHA.items():
        if hashlib.sha256(blobs[name]).hexdigest()!=sha:
            raise ValueError("Unrecognized source XLSX checksum")
    from openpyxl import load_workbook
    sites=[]
    for name,sha in EXPECTED_SHA.items():
        wb=load_workbook(io.BytesIO(blobs[name]),read_only=True,data_only=True)
        try:
            if len(wb.worksheets)!=1:
                raise ValueError("Unexpected original source worksheet count")
            sites.append({"site":name,"sha256":sha,
                          **audit_sheet(wb.worksheets[0],expected_years=ORIGINAL_YEARS)})
        finally:wb.close()
    return {
        "status":"SOURCE_METADATA_AND_MISSINGNESS_AUDITED_NOT_WOFOST_READY",
        "doi":"10.7910/DVN/1LC6W7",
        "source_version":"3.0","source_license":"CC BY-NC-SA 4.0",
        "raw_workbooks_exported":0,
        "daily_weather_values_exported":0,
        "original_header_strings_exported":0,
        "model_input_mapping_approved":False,
        "soil_crop_management_alignment_approved":False,
        "heldout_wofost_predicted":False,
        "field_wofost_run_admissible":False,
        "sites":sites,
        "limits":[
            "Header hints are NOT source-level unit/sensor definitions or PCSE mappings.",
            "Missingness is counted; original blanks are never imputed.",
            "Percent humidity is not VAP hPa without a documented method.",
            "Makkink ET reference is not interchangeable with E0 ES0 and ET0.",
            "Real soil, genotype, management and evaluation holdout remain unresolved.",
            "No Fe/Zn or crop improvement claim can follow from this meteorology.",
        ],
    }

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--lelystad",type=Path,required=True)
    p.add_argument("--vredepeel",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    a=p.parse_args()
    report=audit_originals({
        "Lelystad":a.lelystad.read_bytes(),
        "Vredepeel":a.vredepeel.read_bytes()})
    a.out.mkdir(parents=True,exist_ok=True)
    (a.out/"m1_24_weather_semantics_and_missingness.json").write_text(
        json.dumps(report,sort_keys=True,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({
        "status":report["status"],
        "site_missing_yearly":[{"site":s["site"],"columns":[
            {"column":c["source_column_number"],"years":{
                y: v["missing_na"]+v["untyped_nonfinite"]
                for y,v in c["counts_by_year"].items()}}
            for c in s["columns"] if any(
                v["missing_na"] or v["untyped_nonfinite"]
                for v in c["counts_by_year"].values())]}
            for s in report["sites"]],
        "field_wofost_run_admissible":False},sort_keys=True))

if __name__=="__main__": main()
