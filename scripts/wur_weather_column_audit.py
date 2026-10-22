"""M1-23: original WUR weather numeric coverage, no daily or header exports.

All source data licensed CC BY-NC-SA 4.0, DOI 10.7910/DVN/1LC6W7.
Only positional column counts and coarse pattern flags may leave the runner.
No inference about crop model validity follows from numeric completeness.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
from collections import Counter
from pathlib import Path

from scripts.wur_weather_daily_gate import parse_day

EXPECTED_FILES = {
    "Lelystad": "ad66a23468a969b5795a656fb6989499e790a3e18179b9e3158f203b12ab8636",
    "Vredepeel": "0bb40a5aa91ae6bfd02a8e8b4f825e4d834207884c1f2f44caa12fd0e88a1fd6",
}
METADATA_HINTS = {
    "temperature": r"temperature|temperatuur|temp|tmin|tmax",
    "rainfall": r"precip|rain|neerslag|regen",
    "global_radiation": r"radiat|straling|irrad",
    "humidity": r"humid|vochtig|luchtvocht",
    "vapor_pressure": r"vapou?r\s*press|dampdruk",
    "wind": r"wind",
    "evapotranspiration": r"evap|verdamp|makkink|etref",
    "unit_degC": r"°c|deg(?:rees)?\s*c|celsius",
    "unit_millimeters": r"\bmm(?:/| per | d| day|$)",
    "unit_kilo_joules": r"\bkj\b",
    "unit_percent": r"%",
    "unit_wind_height_2m": r"\b2\s*m\b",
}


def coverage_of_sheet(ws):
    rows = list(ws.iter_rows(min_row=10, values_only=True))
    if not rows:
        raise ValueError("No data rows after original metadata")
    ncol = max(len(r) for r in rows)
    candidates = []
    for col in range(ncol):
        populated = [row[col] for row in rows if col < len(row) and row[col] is not None]
        dates = [parse_day(x) for x in populated]
        valid = [d for d in dates if d is not None and 1950 <= d.year <= 2100]
        if len(valid) >= 100 and len(valid) / len(populated) >= 0.9:
            candidates.append((col,len(valid)))
    if not candidates:
        return {"date_axis_status":"NONE", "columns":[],
                "source_variable_units_verified":False,"ready_for_pcse":False}
    maximum = max(x[1] for x in candidates)
    candidates = [idx for idx,count in candidates if count==maximum]
    if len(candidates)!=1:
        return {"date_axis_status":"AMBIGUOUS","columns":[],
                "source_variable_units_verified":False,"ready_for_pcse":False}
    datecol = candidates[0]
    sample = [(parse_day(row[datecol]),row) for row in rows if len(row)>datecol]
    weather = [(d,r) for d,r in sample if d is not None and d.year in (2019,2020)]
    if len(weather)<100:
        raise ValueError("Insufficient actual 2019-2020 date coverage")
    metarows=list(ws.iter_rows(min_row=1,max_row=9,values_only=True))
    columns=[]
    for idx in range(ncol):
        if idx==datecol:continue
        numeric=missing=untyped=0
        for _,row in weather:
            val=row[idx] if idx<len(row) else None
            if val is None or isinstance(val,str) and val.strip().upper() in ("","NA","NAN","NULL"):
                missing+=1
            elif isinstance(val,bool) or not isinstance(val,(int,float)) or not math.isfinite(val):
                untyped+=1
            else:
                numeric+=1
        meta=" ".join(str(row[idx]) for row in metarows if idx<len(row) and isinstance(row[idx],str)).lower()
        hints=sorted(name for name,pattern in METADATA_HINTS.items()
                     if re.search(pattern,meta,re.IGNORECASE))
        columns.append({
            "source_column_1_based":idx+1,
            "target_2019_2020_calendar_rows":len(weather),
            "finite_numeric_cells":numeric,
            "missing_or_NA_cells":missing,
            "untyped_or_nonfinite_cells":untyped,
            "tentative_header_patterns_not_validated_mapping":hints,
            "units_or_sensor_method_confirmed":False,
        })
    return {
        "date_axis_status":"DETECTED_NOT_PHYSICAL_VARIABLE_MAPPING",
        "2019_2020_calendar_rows_with_date":len(weather),
        "date_axis_source_column_1_based":datecol+1,
        "columns":columns,
        "source_variable_units_verified":False,
        "ready_for_pcse":False,
    }


def inspect_originals(blobs):
    from openpyxl import load_workbook
    if set(blobs)!=set(EXPECTED_FILES):
        raise ValueError("Both source stations required")
    result=[]
    for name,expected in EXPECTED_FILES.items():
        if hashlib.sha256(blobs[name]).hexdigest()!=expected:
            raise ValueError("Original source digest mismatch")
        wb=load_workbook(io.BytesIO(blobs[name]),read_only=True,data_only=True)
        try:
            if len(wb.worksheets)!=1:
                raise ValueError("Original workbook sheet structure needs re-verification")
            details=coverage_of_sheet(wb.worksheets[0])
            result.append({"site":name,"original_sha256":expected,**details})
        finally:
            wb.close()
    return {
        "status":"ORIGINAL_FIELD_WEATHER_NUMERIC_COVERAGE_AUDITED_NOT_MODEL_ADMITTED",
        "source_doi":"10.7910/DVN/1LC6W7",
        "license":"CC BY-NC-SA 4.0",
        "dataset_version":"3.0",
        "original_daily_observation_values_exported":0,
        "original_header_strings_exported":0,
        "weather_series_redistributed":False,
        "source_physical_units_or_methods_approved":False,
        "soil_crop_irrigation_matching_approved":False,
        "field_wofost_run_admissible":False,
        "sites":result,
        "limitations":[
            "Counts and tentative regex hints do NOT validate the published source's meteorological dictionary.",
            "Finite numeric entries do NOT establish correct units, station height, sensor calibration or PCSE compatibility.",
            "An observed percent humidity is NOT directly equivalent to VAP in hPa.",
            "CC BY-NC-SA 4.0 applies to source, no weather values, raw headers or series exported.",
            "No potato yield, Fe/Zn, safety, bioavailability or WOFOST accuracy claims.",
        ],
    }


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--lelystad",type=Path,required=True)
    p.add_argument("--vredepeel",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    a=p.parse_args()
    report=inspect_originals({
        "Lelystad":a.lelystad.read_bytes(),"Vredepeel":a.vredepeel.read_bytes()})
    a.out.mkdir(parents=True,exist_ok=True)
    (a.out/"m1_23_weather_column_numeric_coverage.json").write_text(
        json.dumps(report,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps({"status":report["status"],"field_wofost_run_admissible":False,
        "site_summaries":[{"site":s["site"],"date_axis_status":s["date_axis_status"],
         "columns":s["columns"]} for s in report["sites"]]},sort_keys=True))


if __name__=="__main__":main()
