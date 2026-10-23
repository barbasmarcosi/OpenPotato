"""M1-30: pinned WUR irrigation workbook typed-column & vocabulary audit.
Original ten Den et al. 2024, Harvard Dataverse v3, CC BY-NC-SA 4.0.
Exports coarse counts/whitelisted tags ONLY: never licensed values, dates or rows.
"""
from collections import Counter,defaultdict
from datetime import date,datetime
from pathlib import Path
import argparse,hashlib,io,json,math,re
from scripts.wur_irrigation_source_inventory import EXPECTED_SHA256,parse_date

CUES={
 "date":r"\bdate\b|\bdatum\b|\bday\b",
 "year":r"\byear\b|\bjaar\b",
 "site":r"\blocation\b|\blocatie\b|\bsite\b",
 "rain":r"\brain\b|\bprecipitation\b|\bneerslag\b",
 "irrigation":r"\birrigation\b|\birrigatie\b|\bwatergift\b",
 "mm":r"(?<!\w)mm(?!\w)",
 "treatment":r"\btreatment\b|\bbehandeling\b|\bw1\b|\bw2\b",
 "plot":r"\bplot\b|\bplotnumber\b|\bperceel\b",
}
CATEGORIES={"Lelystad":{"lelystad","l"},"Vredepeel":{"vredepeel","v"},
            "W1":{"w1"},"W2":{"w2"}}

def kind(v):
 if v is None or isinstance(v,str) and v.strip().upper() in ("","NA","NULL","NAN"):return "missing"
 if isinstance(v,(date,datetime)):return "date"
 if isinstance(v,bool):return "boolean"
 if isinstance(v,(int,float)):return "numeric" if math.isfinite(v) else "invalid"
 return "text" if isinstance(v,str) else "invalid"

def whitelisted_tags(rows,limit):
 output={}
 for rownum,row in enumerate(rows[:limit],start=1):
  tags=[]
  for col,value in enumerate(row,start=1):
   if isinstance(value,str):
    recognized=sorted(key for key,pattern in CUES.items() if re.search(pattern,value,re.I))
    if recognized:tags.append({"position":col,"controlled_tags":recognized})
  if tags:output[str(rownum)]=tags
 return output

def verify_site_year_date_keys(rows):
 """Site x year x day event-candidate identities, without source row export."""
 sites={"lelystad":"Lelystad","l":"Lelystad",
        "vredepeel":"Vredepeel","v":"Vredepeel"}
 observed=set()
 totals=Counter()
 numeric=defaultdict(Counter)
 for row in rows[1:]:
  if len(row)!=6 or not isinstance(row[0],str):
   raise ValueError("Malformed source location/site row")
  site=sites.get(row[0].strip().lower())
  if not site:raise ValueError("Unrecognized site label in original irrigation source")
  year=row[1]
  if isinstance(year,bool) or not isinstance(year,(int,float)) or year not in (2019,2020):
   raise ValueError("Unrecognized source year")
  d=parse_date(row[2])
  if d is None or d.year!=int(year):
   raise ValueError("Original year/date mismatch")
  key=(site,int(year),d)
  if key in observed:
   raise ValueError("Duplicate site-year-day row, cannot infer irrigation event")
  observed.add(key)
  group=(site,str(int(year)))
  totals[group]+=1
  for i in (3,4,5):
   v=row[i]
   k=kind(v)
   if k=="numeric" and v<0:raise ValueError("Negative source water-like value")
   if k not in ("numeric","missing"):
    raise ValueError("Unrecognized source water-like column cell")
   numeric[group][(i+1,k)]+=1
 if len(observed)!=526 or len(totals)!=4:
  raise ValueError("Original site-year-day coverage changed")
 return [
  {"site":site,"year":year,"unique_site_year_days":totals[site,year],
   "source_column_numeric_or_missing_counts":{
    str(col):{kind:numeric[site,year][col,kind] for kind in ("numeric","missing")}
    for col in (4,5,6)}}
  for site,year in sorted(totals)
 ]

def _semantic_fingerprints(rows):
 """Publish only boolean matches for a FIXED controlled source vocabulary."""
 rules={
  "mentions_w1":r"(?<![a-z0-9])w\s*1(?![a-z0-9])",
  "mentions_w2":r"(?<![a-z0-9])w\s*2(?![a-z0-9])",
  "mentions_mm":r"(?<![a-z])mm(?![a-z])|millimet",
  "mentions_rain":r"rain|neerslag|precip",
  "mentions_irrigation":r"irrigat|watergift|watering",
  "mentions_pF":r"(?<![a-z])pf\s*[23](?:[.,][24])?",
  "mentions_day":r"date|datum|day",
  "mentions_2019":r"(?<![0-9])2019(?![0-9])",
  "mentions_2020":r"(?<![0-9])2020(?![0-9])",
 }
 result={}
 for rowindex,row in enumerate(rows):
  found=[]
  for colindex,cell in enumerate(row):
   if not isinstance(cell,str):continue
   tags=[k for k,expr in rules.items() if re.search(expr,cell,re.I)]
   if tags:found.append({"column":colindex+1,"matches":sorted(tags)})
  if found:result[str(rowindex+1)]=found
 return result

def inspect_sheets(meta,data):
 metarows=list(meta.iter_rows(values_only=True))
 rows=list(data.iter_rows(values_only=True))
 if len(metarows)!=17 or len(rows)!=527 or data.max_column!=6:
  raise ValueError("Original source workbook geometry changed")
 date_like=Counter()
 for row in rows:
  for col,value in enumerate(row):
   d=parse_date(value)
   if d is not None and d.year in (2019,2020):date_like[col]+=1
 strongest=max(date_like.values(),default=0)
 possible=[col for col,count in date_like.items() if count==strongest and count>=100]
 if len(possible)!=1:raise ValueError("Ambiguous source data date axis")
 date_index=possible[0]
 counts=Counter()
 summaries=[]
 for col in range(6):
  typed=defaultdict(Counter)
  names=Counter()
  for row in rows:
   d=parse_date(row[date_index]) if len(row)>date_index else None
   if d is None or d.year not in (2019,2020):continue
   year=str(d.year)
   if col==date_index:counts[year]+=1
   v=row[col] if col<len(row) else None
   k=kind(v)
   typed[year][k]+=1
   if k=="numeric":
    typed[year]["zero" if v==0 else ("positive" if v>0 else "negative")]+=1
   if isinstance(v,str):
    val=v.strip().lower()
    for category,options in CATEGORIES.items():
     if val in options:names[category]+=1
  summaries.append({"column_one_based":col+1,"date_axis":col==date_index,
                    "counts_by_year":{year:dict(cnt) for year,cnt in sorted(typed.items())},
                    "controlled_site_treatment_category_counts":dict(names),
                    "unit_approved":False,"event_join_approved":False})
 if dict(counts)!={"2019":258,"2020":268}:raise ValueError("Date counts differ from original")
 return {"meta_rows":17,"data_rows":527,"data_columns":6,
         "date_column_one_based":date_index+1,"date_counts_by_year":dict(counts),
         "metadata_tags_by_position":whitelisted_tags(metarows,17),
         "source_metadata_semantic_fingerprints":_semantic_fingerprints(metarows),
         "source_data_header_semantic_fingerprints":_semantic_fingerprints(rows[:1]),
         "data_first_rows_tags_by_position":whitelisted_tags(rows,15),
         "column_profiles":summaries,
         "original_values_or_dates_exported":0,
 "site_year_date_keys_verified":True,
 "site_year_event_candidate_counts":verify_site_year_date_keys(rows)}

def audit(blob):
 if hashlib.sha256(blob).hexdigest()!=EXPECTED_SHA256:
  raise ValueError("Original Irrigation.xlsx SHA256 mismatch")
 from openpyxl import load_workbook
 workbook=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
 try:
  sheets={sheet.title.strip().lower():sheet for sheet in workbook.worksheets}
  if set(sheets)!={"meta","data"}:raise ValueError("Original irrigation sheets changed")
  profile=inspect_sheets(sheets["meta"],sheets["data"])
 finally:workbook.close()
 return {"status":"M1_30_ORIGINAL_IRRIGATION_SOURCE_PROFILE_NO_AGROMANAGEMENT",
         "doi":"10.7910/DVN/1LC6W7","file_id":10284791,
         "source_sha256":EXPECTED_SHA256,"license":"CC BY-NC-SA 4.0",
         "schema_hints_are_not_approved_field_mappings":True,
         "water_units_validated":False,"site_plot_treatment_join_verified":False,
         "2019_train_2020_evaluation_preserved":True,
         "wofost_field_run_admissible":False,"original_records_exported":0,
         "profile":profile}

def main():
 p=argparse.ArgumentParser(description=__doc__)
 p.add_argument("--source",type=Path,required=True)
 p.add_argument("--out",type=Path,required=True)
 a=p.parse_args()
 result=audit(a.source.read_bytes())
 a.out.mkdir(parents=True,exist_ok=True)
 (a.out/"m1_30_irrigation_source_column_profile.json").write_text(
  json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
 print(json.dumps({"status":result["status"],"profile":result["profile"],
                   "wofost_ready":False},sort_keys=True))
if __name__=="__main__":main()
