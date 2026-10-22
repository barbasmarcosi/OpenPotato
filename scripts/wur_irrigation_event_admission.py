"""M1-31: source-bound, read-only WUR irrigation treatment identity gate.

ten Den et al. 2024 Harvard Dataverse v3 CC BY-NC-SA 4.0.
Checks the ORIGINAL source metadata, weathering treatment allocations and
event-candidate integrity. No original source measurements, plot IDs or
per-day records leave the process. NO WOFOST agromanagement is emitted.
"""
import argparse,csv,hashlib,io,json,math,re
from collections import Counter,defaultdict
from datetime import date
from pathlib import Path
from scripts.wur_irrigation_source_inventory import EXPECTED_SHA256,parse_date
from scripts import wur_potato_field_audit as plotsource

SITE_CODES={"lelystad":"L","vredepeel":"V"}
SITE_NAMES={"L":"Lelystad","V":"Vredepeel"}
YEARS=(2019,2020)
REQUIRED_COLUMNS={"rainfall":3,"W1":4,"W2":5} # zero-based physical positions
FIELD_SIGNATURE={
 "rainfall":r"\brain|neerslag|precipitation",
 "W1":r"irrigat|watergift|waterapp",
 "W2":r"irrigat|watergift|waterapp",
}
W_CODES={"W1":r"(?<!\w)w\s*1(?!\w)","W2":r"(?<!\w)w\s*2(?!\w)"}
UNIT_PATTERNS={
 "mm":r"(?<![a-z])mm(?![a-z])|\bmillimet(?:re|er)",
 "cm":r"(?<![a-z])cm(?![a-z])|\bcentimet(?:re|er)",
 "litres_per_m2":r"\bl(?:itres?|iters?)?\s*/\s*m(?:2|²)\b",
 "other":r"\bm3\b|\bliters?\b|\blitres?\b",
}

def safe_source_signatures(meta,data):
 """Each positional mapping must be named independently in DATA and META."""
 head=next(data.iter_rows(min_row=1,max_row=1,values_only=True),None)
 if head is None or len(head)!=6:raise ValueError("Original irrigation header length changed")
 rows=list(meta.iter_rows(values_only=True))
 if len(rows)!=17:raise ValueError("Irrigation source metadata row count changed")
 expected=[("rainfall",14),("W1",15),("W2",16)]
 matches={}
 for name,meta_idx in expected:
  col=REQUIRED_COLUMNS[name]
  h=head[col]
  row=rows[meta_idx]
  if not isinstance(h,str) or not re.search(FIELD_SIGNATURE[name],h,re.I):
   raise ValueError("Irrigation data header not source-validated for "+name)
  if name!="rainfall" and not re.search(W_CODES[name],h,re.I):
   raise ValueError("Irrigation treatment code absent from data header "+name)
  if len(row)<3 or not isinstance(row[0],str) or not isinstance(row[2],str):
   raise ValueError("Original treatment column dictionary missing")
  for source_label in (row[0],row[2]):
   if not re.search(FIELD_SIGNATURE[name],source_label,re.I):
    raise ValueError("Irrigation source metadata label differs for "+name)
   if name!="rainfall" and not re.search(W_CODES[name],source_label,re.I):
    raise ValueError("Metadata treatment code mismatch "+name)
  matches[name]={"data_column_one_based":col+1,
                 "source_data_label_recognized":True,
                 "source_metadata_row_one_based":meta_idx+1,
                 "source_metadata_label_recognized":True}
 # Detection only; there may be no unit in the published XLSX dictionary.
 source_meta=" ".join(str(v) for row in rows for v in row if isinstance(v,str))
 source_headers=" ".join(v for v in head if isinstance(v,str))
 units=sorted(k for k,pattern in UNIT_PATTERNS.items()
              if re.search(pattern,source_meta+" "+source_headers,re.I))
 unit_status=("explicit_mm" if units==["mm"] else
              "explicit_litres_per_m2" if units==["litres_per_m2"] else
              "AMBIGUOUS_OR_UNSPECIFIED")
 return matches,unit_status,units

def inspect_event_candidates(rows):
 """Never make zero out of a missing W1/W2 source cell."""
 seen=set();aggregates=defaultdict(lambda:defaultdict(Counter))
 for row in rows[1:]:
  if len(row)!=6:raise ValueError("Irrigation row physical width changed")
  site=SITE_CODES.get(str(row[0]).strip().lower())
  if site is None:raise ValueError("Original irrigation site unknown")
  year=row[1]
  if isinstance(year,bool) or not isinstance(year,(float,int)) or int(year)!=year or int(year) not in YEARS:
   raise ValueError("Source irrigation year malformed")
  year=int(year)
  d=parse_date(row[2])
  if d is None or d.year!=year:raise ValueError("Date/year mismatch")
  key=(site,year,d)
  if key in seen:raise ValueError("Duplicate irrigation site-year-day")
  seen.add(key)
  group=(year,site)
  aggregates[group]["calendar"]["dated_rows"]+=1
  for kind,col in REQUIRED_COLUMNS.items():
   val=row[col]
   if val is None or isinstance(val,str) and val.strip().upper() in ("","NA","NULL","NAN"):
    if kind=="rainfall":raise ValueError("Rainfall missing, not a zero")
    aggregates[group][kind]["source_missing_not_zero"]+=1
    continue
   if isinstance(val,bool) or not isinstance(val,(float,int)) or not math.isfinite(val) or val<0:
    raise ValueError("Invalid or negative source water field")
   aggregates[group][kind]["observed_numeric"]+=1
   if val==0:aggregates[group][kind]["observed_explicit_zero"]+=1
   else:aggregates[group][kind]["observed_positive"]+=1
 if len(seen)!=526 or len(aggregates)!=4:
  raise ValueError("Pinned original source site/day counts changed")
 return aggregates

def inspect_plot_assignments(blob):
 """ONLY irrigation codes and plot IDs are accessed. Biomass is untouched."""
 plotsource.inspect_original(blob)
 reader=csv.DictReader(io.StringIO(blob.decode("utf-8-sig")))
 observed={}
 for row in reader:
  year=row["Year"].strip()
  site=row["Location"].strip()
  plot=row["Plotnumber"].strip()
  trt=row["Irrigation"].strip().upper()
  if site not in SITE_NAMES or year not in ("2019","2020") or not plot:
   raise ValueError("Experimental plot identity outside recognized study")
  if trt not in W_CODES:
   raise ValueError("Unmatched plot irrigation treatment CODE; do not guess mapping")
  key=(int(year),site,plot)
  previous=observed.setdefault(key,trt)
  if previous!=trt:raise ValueError("Within-plot irrigation assignment changed")
 if len(observed)!=174:raise ValueError("Expected physical plot-year identities changed")
 n=defaultdict(Counter)
 for (year,site,plot),trt in observed.items():
  n[(year,site)][trt]+=1
 for year in YEARS:
  for site,expected in (("L",45),("V",42)):
   if sum(n[year,site].values())!=expected or not all(n[year,site][w]>0 for w in ("W1","W2")):
    raise ValueError("No balanced real plot treatment join")
 return n

def audit(blob_irrigation,blob_plots):
 if hashlib.sha256(blob_irrigation).hexdigest()!=EXPECTED_SHA256:
  raise ValueError("Pinned original irrigation SHA256 mismatch")
 if hashlib.sha256(blob_plots).hexdigest()!=plotsource.SOURCE_SHA256:
  raise ValueError("Pinned original plot CSV SHA256 mismatch")
 from openpyxl import load_workbook
 book=load_workbook(io.BytesIO(blob_irrigation),read_only=True,data_only=True)
 try:
  sheets={s.title.lower():s for s in book.worksheets}
  if set(sheets)!={"meta","data"}:raise ValueError("Unexpected irrigation source sheets")
  matches,units_status,unit_clues=safe_source_signatures(sheets["meta"],sheets["data"])
  rows=list(sheets["data"].iter_rows(values_only=True))
  events=inspect_event_candidates(rows)
 finally:book.close()
 plots=inspect_plot_assignments(blob_plots)
 report=[]
 for year,site in sorted(events):
  a=events[year,site]
  report.append({
   "site":SITE_NAMES[site],"year":year,
   "recorded_unique_site_year_days":a["calendar"]["dated_rows"],
   "rainfall_positive_days":a["rainfall"]["observed_positive"],
   "rainfall_explicit_zero_days":a["rainfall"]["observed_explicit_zero"],
   "by_treatment":{
    name:{"recorded_positive_water_amount_cells":a[name]["observed_positive"],
          "observed_zero_amount_cells":a[name]["observed_explicit_zero"],
          "unrecorded_blank_cells_not_zero":a[name]["source_missing_not_zero"],
          "matching_original_plot_assignments":plots[year,site][name]}
    for name in ("W1","W2")},
  })
 return {
  "status":"M1_31_ORIGINAL_WUR_RAIN_W1_W2_HEADERS_AND_PLOT_CODES_VERIFIED",
  "source_licence":"CC BY-NC-SA 4.0",
  "irrigation_file_sha256":EXPECTED_SHA256,
  "plots_file_sha256":plotsource.SOURCE_SHA256,
  "source_columns":matches,
  "unit_label_status":units_status,
  "controlled_source_unit_clues":unit_clues,
  "all_site_year_treatment_labels_reconciled":True,
  "site_year_aggregated_counts":report,
  "raw_source_rows_or_irrigation_doses_exported":0,
  "plot_ids_exported":0,"raw_individual_dates_exported":0,
  "irrigation_mm_per_day_approved":units_status in ("explicit_mm","explicit_litres_per_m2"),
  "complete_zero_imputation_permitted":False,
  "2020_field_biomass_used_for_calibration":False,
  "wofost_field_simulation_validated":False,
  "wofost_agromanagement_ready":False,
  "limitations":[
   "A recorded positive water value is not a full continuous zero-filled irrigation schedule.",
   "Treat unrecorded W1/W2 cells as UNKNOWN, never automatic zero.",
   "A valid site-year-treatment association does not validate the site's soil/cultivar model.",
   "Published rainfall and weather files can be compared in a future source-quality gate.",
   "2019 training and 2020 reserved validation remain distinct; no crop model run.",
  ],
 }

def main():
 p=argparse.ArgumentParser(description=__doc__)
 p.add_argument("--irrigation",required=True,type=Path)
 p.add_argument("--plots",required=True,type=Path)
 p.add_argument("--out",required=True,type=Path)
 a=p.parse_args()
 r=audit(a.irrigation.read_bytes(),a.plots.read_bytes())
 a.out.mkdir(parents=True,exist_ok=True)
 (a.out/"m1_31_source_verified_treatment_event_counts_only.json").write_text(
  json.dumps(r,indent=2,sort_keys=True)+"\n",encoding="utf-8")
 print(json.dumps({"status":r["status"],"units":r["unit_label_status"],
                   "source_columns":r["source_columns"],
                   "site_year_counts":r["site_year_aggregated_counts"],
                   "agromanagement_ready":False},sort_keys=True))
if __name__=="__main__":main()
