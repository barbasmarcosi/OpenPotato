"""M1-32 — Cross-source WUR rainfall audit (licensed observations in memory only).

Harvard Dataverse ten Den et al. v3.0 CC BY-NC-SA 4.0:
Irrigation.xlsx versus hash-pinned Weatherfile_{lelystad,vredepeel}.xlsx.

Compare RAIN reported in native weather input mm to the named daily 'rainfall'
column in irrigation register. Cross-source *agreement is not calibration*,
proof of irrigation units, blank=zero or PCSE agromanagement validation.
Release ONLY aggregate counts / rounded group comparison metrics.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import hashlib
import io
import json
import math
from pathlib import Path

from scripts.wur_irrigation_source_inventory import EXPECTED_SHA256, parse_date
from scripts.wur_irrigation_event_admission import (
    SITE_CODES, SITE_NAMES, REQUIRED_COLUMNS, safe_source_signatures,
)
from scripts.wur_pcse_source_contract import HASHES

EXPECTED_DATE_COUNTS={(2019,"L"):127,(2019,"V"):131,
                      (2020,"L"):134,(2020,"V"):134}

def valid_number(value, label):
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):
        raise ValueError("Missing or non-finite "+label)
    if value<0:raise ValueError("Negative "+label)
    return float(value)

def index_irrigation(meta, data):
    """Build a PRIVATE in-memory pairing key and typed water-cell availability."""
    fields, units_status, _ = safe_source_signatures(meta,data)
    if {k:v["data_column_one_based"] for k,v in fields.items()} != {
        "rainfall":4,"W1":5,"W2":6}:
        raise ValueError("Irrigation source field positions changed")
    rows=data.iter_rows(min_row=1,values_only=True)
    header=next(rows,None)
    if header is None or len(header)!=6:raise ValueError("Missing irrigation source header")
    entries={}
    counts=Counter()
    for row in rows:
        if len(row)!=6:raise ValueError("Malformed irrigation row width")
        if not isinstance(row[0],str):
            raise ValueError("Unknown location code")
        site=SITE_CODES.get(row[0].strip().lower())
        if site is None:raise ValueError("Unknown irrigation source site")
        year_raw=row[1]
        if isinstance(year_raw,bool) or not isinstance(year_raw,(int,float)) or year_raw not in (2019,2020):
            raise ValueError("Irrigation year malformed")
        year=int(year_raw)
        d=parse_date(row[2])
        if d is None or d.year!=year:raise ValueError("Irrigation year/date inconsistent")
        key=(year,site,d)
        if key in entries:raise ValueError("Duplicate irrigation site/year/day")
        rain=valid_number(row[REQUIRED_COLUMNS["rainfall"]],"irrigation-record rainfall")
        status={}
        for trt in ("W1","W2"):
            value=row[REQUIRED_COLUMNS[trt]]
            if value is None or isinstance(value,str) and value.strip().upper() in ("","NA","NULL","NAN"):
                status[trt]="unknown_blank"
            else:
                number=valid_number(value,"source "+trt+" irrigation")
                status[trt]="explicit_zero" if number==0 else "recorded_positive"
        entries[key]=(rain,status)
        counts[year,site]+=1
    if counts!=Counter(EXPECTED_DATE_COUNTS):
        raise ValueError("Original irrigation site-year row balance changed")
    return entries, units_status

def index_weather(ws,site):
    """Index site-specific original daily RAIN as millimeters (source declared)."""
    rows=ws.iter_rows(min_row=1,values_only=True)
    header_row=None
    unit_row=None
    indexed={}
    for line,row in enumerate(rows,start=1):
        if line==10:header_row=row
        elif line==11:unit_row=row
        elif line>=12:
            if header_row is None or unit_row is None:
                raise ValueError("Original weather labels/units missing")
            date_i=header_row.index("DAY")
            rain_i=header_row.index("RAIN")
            d=parse_date(row[date_i])
            if d is None:raise ValueError("Unparseable original weather date")
            if d in indexed:raise ValueError("Duplicate original weather date")
            indexed[d]=valid_number(row[rain_i],"source weather RAIN")
    if header_row is None or unit_row is None:
        raise ValueError("Original weather schema missing")
    if sum(x=="RAIN" for x in header_row)!=1 or sum(x=="DAY" for x in header_row)!=1:
        raise ValueError("Weather DAY/RAIN labels ambiguous")
    rain_i=header_row.index("RAIN")
    if unit_row[rain_i] != "mm":
        raise ValueError("Weather RAIN source declared unit is not mm")
    if len(indexed)!=1096:
        raise ValueError("Unexpected original weather daily coverage")
    return indexed

def compare_indices(irrigation,weather_by_site,irrigation_unit_status):
    """No source-specific daily values or dates leave these aggregates."""
    accum=defaultdict(lambda:{"n":0,"exact":0,"tenfold_equal":0,"tenfold_mismatch":0,"tenfold_max_residual":0.,"both_zero":0,
       "both_positive":0,"register_positive_weather_zero":0,
       "weather_positive_register_zero":0,"numerical_abs_difference":0.,
       "rainfall_register_sum":0.,"rainfall_weather_sum":0.,
       "w1":Counter(),"w2":Counter()})
    for (year,site,d),(rain,status) in irrigation.items():
        if site not in weather_by_site or d not in weather_by_site[site]:
            raise ValueError("Missing source weather day for irrigation register")
        weather_rain=weather_by_site[site][d]
        group=accum[year,site]
        group["n"]+=1
        if abs(rain-weather_rain) <= 1e-9:group["exact"]+=1
        # An empirical factor-of-ten check, not a unit/sensor certification.
        residual=abs(10*rain-weather_rain)
        group["tenfold_max_residual"]=max(group["tenfold_max_residual"],residual)
        if residual<=1e-7:group["tenfold_equal"]+=1
        else:group["tenfold_mismatch"]+=1
        if rain==0 and weather_rain==0:group["both_zero"]+=1
        elif rain>0 and weather_rain>0:group["both_positive"]+=1
        elif rain>0:group["register_positive_weather_zero"]+=1
        else:group["weather_positive_register_zero"]+=1
        group["numerical_abs_difference"]+=abs(rain-weather_rain)
        group["rainfall_register_sum"]+=rain
        group["rainfall_weather_sum"]+=weather_rain
        for name in ("W1","W2"):
            group[name.lower()][status[name]]+=1
    if set(accum)!=set(EXPECTED_DATE_COUNTS):raise ValueError("Missing site-year group")
    result=[]
    for year,site in sorted(accum):
        v=accum[year,site]
        if v["n"]!=EXPECTED_DATE_COUNTS[year,site]:
            raise ValueError("Original join coverage changed")
        rain_classified=(v["both_zero"]+v["both_positive"]+
              v["register_positive_weather_zero"]+
              v["weather_positive_register_zero"])
        if rain_classified!=v["n"]:raise ValueError("Rain class counts inconsistent")
        for name in ("W1","W2"):
            if sum(v[name.lower()].values())!=v["n"]:
                raise ValueError("Irrigation treatment presence counts inconsistent")
        result.append({
            "year":year,"site":SITE_NAMES[site],
            "same_site_same_day_records":v["n"],
            "rain_positive_in_both_sources_days":v["both_positive"],
            "rain_zero_in_both_sources_days":v["both_zero"],
            "rain_daily_values_matching_tenfold_weather_mm_relation":v["tenfold_equal"],
            "rain_daily_values_NOT_matching_tenfold_relation":v["tenfold_mismatch"],
            "max_absolute_tenfold_residual_in_weather_numeric_scale":round(v["tenfold_max_residual"],8),
            "irrigation_register_rain_positive_weather_zero_days":v["register_positive_weather_zero"],
            "weather_rain_positive_irrigation_register_zero_days":v["weather_positive_register_zero"],
            "equal_rainfall_numerical_cells_if_both_interpreted_as_mm":v["exact"],
            "conditional_MAE_in_original_numeric_scales":round(v["numerical_abs_difference"]/v["n"],3),
            "conditional_rainfall_totals_ratio_register_over_weather":(
               round(v["rainfall_register_sum"]/v["rainfall_weather_sum"],3)
               if v["rainfall_weather_sum"]>0 else None),
            "treatment_source_cell_status_counts":{
               name:dict(sorted(v[name.lower()].items())) for name in ("W1","W2")
            },
        })
    return {
      "status":"M1_32_CROSS_SOURCE_RAINFALL_QC_NO_IRRIGATION_UNIT_ADMISSION",
      "source_doi":"10.7910/DVN/1LC6W7",
      "source_version":"3.0",
      "source_license":"CC BY-NC-SA 4.0",
      "source_sha256":{
         "Irrigation.xlsx":EXPECTED_SHA256,
         "Weatherfile_lelystad.xlsx":HASHES["Lelystad"],
         "Weatherfile_vredepeel.xlsx":HASHES["Vredepeel"],
      },
      "weather_rain_source_declared_unit":"mm",
      "irrigation_register_water_unit_approved":False,
      "irrigation_register_header_unit_status":irrigation_unit_status,
      "all_526_paired_rainfall_days_observe_10x_source_numeric_relation":all(
          x["rain_daily_values_matching_tenfold_weather_mm_relation"]==x["same_site_same_day_records"]
          for x in result),
      "RAIN_field_independently_declared_mm_but_register_declared_cm":False,
      "rainfall_inter_source_agreement_is_NOT_an_unit_conversion":True,
      "source_humidity_or_rain_sensor_equivalence_approved":False,
      "blanks_treated_as_zero":False,
      "irrigation_W1_W2_mm_approved":False,
      "site_crop_soil_agromanagement_ready":False,
      "2020_crop_biomass_consulted":False,
      "WOFOST_crop_predictions_performed":False,
      "original_daily_dates_or_measurements_exported":0,
      "site_year_aggregates":result,
      "limitations":[
        "The two original rainfall files may represent different sensors or measurement protocols.",
        "Even exact 10x numerical equality does not independently validate W1/W2 source units.",
        "MAE and ratio are numerical diagnostics; neither establishes shared physical units.",
        "Nothing here proves the scale of positive W1/W2 applications in millimeters.",
        "Unknown blanks in W1/W2 have not been turned into zero or no-irrigation events.",
        "No WOFOST agromanagement or cultivar-specific simulation has been admitted.",
      ],
    }

def audit_originals(books):
    if set(books)!={"Irrigation.xlsx","Lelystad","Vredepeel"}:
        raise ValueError("All three original rainfall source files required")
    if hashlib.sha256(books["Irrigation.xlsx"]).hexdigest()!=EXPECTED_SHA256:
        raise ValueError("Irrigation original SHA256 mismatch")
    for site,sha in HASHES.items():
        if hashlib.sha256(books[site]).hexdigest()!=sha:
            raise ValueError("Original "+site+" weather SHA256 mismatch")
    from openpyxl import load_workbook
    irrigation_wb=load_workbook(io.BytesIO(books["Irrigation.xlsx"]),read_only=True,data_only=True)
    try:
        sheets={s.title.strip().lower():s for s in irrigation_wb.worksheets}
        if set(sheets)!={"meta","data"}:raise ValueError("Unexpected irrigation workbook sheets")
        irrigation,units=index_irrigation(sheets["meta"],sheets["data"])
    finally:irrigation_wb.close()
    weather={}
    for site in HASHES:
        wb=load_workbook(io.BytesIO(books[site]),read_only=True,data_only=True)
        try:
            if len(wb.worksheets)!=1:raise ValueError("Unexpected weather workbook sheets")
            weather["L" if site=="Lelystad" else "V"]=index_weather(wb.worksheets[0],site)
        finally:wb.close()
    return compare_indices(irrigation,weather,units)

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--irrigation",type=Path,required=True)
    parser.add_argument("--lelystad",type=Path,required=True)
    parser.add_argument("--vredepeel",type=Path,required=True)
    parser.add_argument("--out",type=Path,required=True)
    args=parser.parse_args()
    report=audit_originals({
       "Irrigation.xlsx":args.irrigation.read_bytes(),
       "Lelystad":args.lelystad.read_bytes(),
       "Vredepeel":args.vredepeel.read_bytes(),
    })
    args.out.mkdir(parents=True,exist_ok=True)
    (args.out/"m1_32_cross_source_rainfall_counts_only.json").write_text(
        json.dumps(report,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps({
      "status":report["status"],
      "units_approved":report["irrigation_W1_W2_mm_approved"],
      "site_year_rainfall_count_diagnostics":report["site_year_aggregates"],
      "WOFOST_ready":False,
    },sort_keys=True))

if __name__=="__main__":
    main()
