"""M1-26 source-bound PCSE input schema gate (NOT crop-model validation).
Original ten Den et al. v3.0 CC BY-NC-SA 4.0. No original data exported.
"""
import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import io
import json
import math
from pathlib import Path

HASHES={
 "Lelystad":"ad66a23468a969b5795a656fb6989499e790a3e18179b9e3158f203b12ab8636",
 "Vredepeel":"0bb40a5aa91ae6bfd02a8e8b4f825e4d834207884c1f2f44caa12fd0e88a1fd6",
}
FIELDS=("DAY","IRRAD","TMIN","TMAX","VAP","WIND","RAIN","SNOWDEPTH")
UNITS=("date","kJ/m2/day or hours","Celsius","Celsius","kPa","m/sec","mm","cm")
DAYS={2018:365,2019:365,2020:366}
MANDATORY=FIELDS[1:-1]

def day(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    return None

def is_missing(value,sentinel):
    if value is None or isinstance(value,str) and value.strip().upper() in ("","NA","NAN","NULL"):
        return True
    if isinstance(value,bool):return True
    try:v=float(value)
    except (ValueError,TypeError):return True
    return not math.isfinite(v) or abs(v-sentinel)<1e-4

def sunshine_flag(value):
    if isinstance(value,bool):return value
    if type(value)==int and value in (0,1):return bool(value)
    if isinstance(value,str):
        v=value.lower().strip()
        if v in ("true","=true","=true()"):return True
        if v in ("false","=false","=false()"):return False
    raise ValueError("Unknown solar radiation versus sunshine flag")

def audit_sheet(ws):
    rows=list(ws.iter_rows(min_row=1,values_only=True))
    if len(rows)<100:raise ValueError("Source weather workbook unexpectedly small")
    # The upstream notebook selects fields BY NAME from row 10, not position.
    # Recognize known PCSE labels only; never export unmatched original strings.
    if len(rows)<12:raise ValueError("Weather header not found")
    labels=rows[9]
    indexes={}
    for field in FIELDS:
        matches=[i for i,v in enumerate(labels)
                 if isinstance(v,str) and v.strip()==field]
        if len(matches)!=1:
            known={f:[i+1 for i,v in enumerate(labels)
                      if isinstance(v,str) and v.strip()==f] for f in FIELDS}
            raise ValueError("Missing or duplicate required PCSE weather field: "+
                             field+"; whitelisted positions: "+str(known))
        indexes[field]=matches[0]
    expected_units={f:u.lower().replace(" ","") for f,u in zip(FIELDS,UNITS)}
    matches={f:isinstance(rows[10][indexes[f]],str) and
                 rows[10][indexes[f]].lower().replace(" ","")==expected_units[f]
             for f in FIELDS}
    if not all(matches.values()):
        raise ValueError("Weather input units disagree for whitelisted fields: "+
                         str([f for f,okay in matches.items() if not okay]))
    try:
        metadata=rows[7]
        lon,lat,elev,a,b=(float(metadata[i]) for i in range(5))
        sunshine=sunshine_flag(metadata[5])
        sentinel=float(rows[5][1])
    except (TypeError,ValueError,IndexError) as exc:
        raise ValueError("Invalid station header needed for PCSE") from exc
    if not (-180<=lon<=180 and -90<=lat<=90 and -500<=elev<=9000):
        raise ValueError("Unphysical site coordinates")
    if not (0<=a<=1 and 0<=b<=1 and 0<a+b<=1):
        raise ValueError("Invalid Angstrom coefficients")
    measured=[]
    for record in rows[11:]:
        if not record:continue
        d=day(record[indexes["DAY"]])
        if d is None:raise ValueError("Unknown day in source series")
        measured.append((d,record))
    dates=[d for d,_ in measured]
    if dates!=sorted(dates) or len(dates)!=len(set(dates)):
        raise ValueError("Duplicate/out of order source dates")
    if (dates[-1]-dates[0]).days+1!=len(dates):
        raise ValueError("Gaps in original source weather chronology")
    if dict(Counter(d.year for d in dates))!=DAYS:
        raise ValueError("Original year coverage changed")
    if max(map(len,rows))!=10:raise ValueError("Original physical weather columns changed")
    essential={f:{"2019":Counter(),"2020":Counter()} for f in MANDATORY}
    unused=[i for i in range(10) if i not in set(indexes.values())]
    extra={i+1:{"2019":Counter(),"2020":Counter()} for i in unused}
    snow={"2019":Counter(),"2020":Counter()}
    for d,rec in measured:
        if d.year not in (2019,2020):continue
        y=str(d.year)
        for name in FIELDS[1:]:
            idx=indexes[name]
            x=rec[idx] if idx<len(rec) else None
            missing=is_missing(x,sentinel)
            target=snow if name=="SNOWDEPTH" else essential[name]
            target[y]["missing" if missing else "numeric"]+=1
            if not missing and name in MANDATORY:
                v=float(x)
                upper={"IRRAD":24 if sunshine else 40000,"TMIN":65,"TMAX":65,
                       "VAP":25,"WIND":50,"RAIN":2500}[name]
                lower=-65 if name in ("TMIN","TMAX") else 0
                if not lower<=v<=upper:
                    raise ValueError("Original mandatory weather outside physical range")
        for idx in unused:
            x=rec[idx] if idx<len(rec) else None
            extra[idx+1][y]["missing" if is_missing(x,sentinel) else "numeric"]+=1
    def dump(counter):
        return {k:{yr:dict(v) for yr,v in years.items()} for k,years in counter.items()}
    mandatory_ok=all(
        counts[year]["numeric"]==DAYS[int(year)]
        for counts in essential.values() for year in ("2019","2020"))
    if not mandatory_ok:raise ValueError("Mandatory weather missing; PCSE may silently skip days")
    return {
        "schema_exact_native_pcse":True,
        "source_units_exact_pcse_excel_input":True,
        "pcse_source_columns_one_based":{f:indexes[f]+1 for f in FIELDS},
        "unused_source_columns_one_based":[i+1 for i in unused],
        "sunshine_hours_mode_from_source_header":sunshine,
        "original_daily_count":len(measured),
        "mandatory_2019_2020_complete":mandatory_ok,
        "mandatory_field_coverage":dump(essential),
        "optional_snow_coverage":{y:dict(c) for y,c in snow.items()},
        "unused_extra_field_coverage":dump(extra),
        "weather_observation_values_exported":0,
        "original_header_strings_exported":0,
        "wind_height_independently_verified":False,
        "E0_ES0_ET0_reference_method_field_validated":False,
        "wofost_crop_model_field_validated":False,
    }

def audit_sources(data):
    if set(data)!=set(HASHES):raise ValueError("Both original source stations required")
    for site,digest in HASHES.items():
        if hashlib.sha256(data[site]).hexdigest()!=digest:
            raise ValueError("Original WUR weather SHA256 mismatch")
    from openpyxl import load_workbook
    sites=[]
    for site in HASHES:
        wb=load_workbook(io.BytesIO(data[site]),read_only=True,data_only=True)
        try:
            if len(wb.worksheets)!=1:raise ValueError("Unexpected workbook sheets")
            sites.append({"site":site,"sha256":HASHES[site],**audit_sheet(wb.worksheets[0])})
        finally:wb.close()
    return {
        "status":"M1_26_ORIGINAL_PCSE_NATIVE_FIELDS_AUDITED_NO_MODEL_VALIDATION",
        "source_doi":"10.7910/DVN/1LC6W7","source_license":"CC BY-NC-SA 4.0",
        "source_version":"3.0",
        "reference_adapter":"ronvree/diffwofost-tutorial (EUPL-1.1); schema inspected, code not copied",
        "PCSE_expected_reader":"ExcelWeatherDataProvider",
        "original_source_records_exported":0,
        "wofost_crop_model_field_validated":False,
        "sites":sites,
        "limitations":[
            "The first eight physical columns are PCSE weather inputs; 9 and 10 are not used.",
            "PCSE converts kJ->J, kPa->hPa, mm->cm and derives E0,ES0,ET0.",
            "Schema and completeness do not validate sensor height or agronomic crop model.",
            "2019 calibration/2020 holdout remains reserved; no model results evaluated.",
        ],
    }

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--lelystad",type=Path,required=True)
    p.add_argument("--vredepeel",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    args=p.parse_args()
    result=audit_sources({"Lelystad":args.lelystad.read_bytes(),
                          "Vredepeel":args.vredepeel.read_bytes()})
    args.out.mkdir(parents=True,exist_ok=True)
    (args.out/"m1_26_pcse_source_contract_counts.json").write_text(
        json.dumps(result,sort_keys=True,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({
        "status":result["status"],"fields":[{"site":s["site"],
          "mandatory_complete":s["mandatory_2019_2020_complete"],
          "extra_columns":s["unused_extra_field_coverage"]}
          for s in result["sites"]],
        "wofost_validated":False},sort_keys=True))

if __name__=="__main__":main()
