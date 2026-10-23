"""M1-33: plot-independent original WUR irrigation window coverage.

Original ten Den et al Harvard Dataverse v3.0 CC BY-NC-SA 4.0.
No raw source doses, individual dates, plots or daily series are exported.
Missing rows and blank irrigation cells are UNKNOWN, not zero.
"""
import argparse
from collections import Counter, defaultdict
from datetime import date
import hashlib
import io
import json
import math
from pathlib import Path

from scripts.wur_agromanagement_admission import PLANT, LAST_SAMPLE
from scripts.wur_irrigation_source_inventory import EXPECTED_SHA256, parse_date
from scripts.wur_irrigation_event_admission import SITE_CODES, SITE_NAMES, safe_source_signatures
from scripts.wur_rainfall_cross_source_qc import EXPECTED_DATE_COUNTS

STATES=("recorded_positive","explicit_zero","blank_UNKNOWN")
WINDOWS=("before_planting","planting_through_last_sample","after_last_sample")

def water_state(value):
    if value is None or isinstance(value,str) and value.strip().upper() in ("","NA","NAN","NULL"):
        return "blank_UNKNOWN"
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0:
        raise ValueError("Malformed source water measurement")
    return "explicit_zero" if value==0 else "recorded_positive"

def assess_rows(rows, *, strict=True):
    seen=set()
    counts=Counter()
    groups=defaultdict(lambda:defaultdict(lambda:defaultdict(Counter)))
    for row in rows:
        if len(row)!=6:raise ValueError("Unexpected source row width")
        if not isinstance(row[0],str):raise ValueError("Source site identifier missing")
        site=SITE_CODES.get(row[0].strip().lower())
        year=row[1]
        if site is None or isinstance(year,bool) or not isinstance(year,(float,int)) or year not in (2019,2020):
            raise ValueError("Unknown original site/year")
        year=int(year)
        d=parse_date(row[2])
        if d is None or d.year!=year:raise ValueError("Source year/date mismatch")
        key=(year,site,d)
        if key in seen:raise ValueError("Repeated irrigation site/year/day")
        seen.add(key)
        start,end=PLANT[str(year),site],LAST_SAMPLE[str(year),site]
        window="before_planting" if d<start else ("after_last_sample" if d>end else "planting_through_last_sample")
        status=water_state(row[3])
        if status=="blank_UNKNOWN":raise ValueError("Source rainfall missing")
        groups[year,site][window]["rain"][status]+=1
        for i,field in ((4,"W1"),(5,"W2")):
            groups[year,site][window][field][water_state(row[i])]+=1
        counts[year,site]+=1
    if set(counts)!=set(EXPECTED_DATE_COUNTS):
        raise ValueError("Four experimental site/year groups required")
    if strict and counts!=Counter(EXPECTED_DATE_COUNTS):
        raise ValueError("Original source record counts changed")
    result=[]
    for year,site in sorted(counts):
        start,end=PLANT[str(year),site],LAST_SAMPLE[str(year),site]
        calendar=(end-start).days+1
        summaries={}
        for win in WINDOWS:
            entry=groups[year,site][win]
            days=sum(entry["rain"].values())
            if any(sum(entry[x].values())!=days for x in ("W1","W2")):
                raise ValueError("Unbalanced treatment availability")
            summaries[win]={
                "recorded_days":days,
                "rain_positive_days":entry["rain"]["recorded_positive"],
                "rain_explicit_zero_days":entry["rain"]["explicit_zero"],
                "W1":{status:entry["W1"][status] for status in STATES},
                "W2":{status:entry["W2"][status] for status in STATES},
            }
        during=summaries["planting_through_last_sample"]
        if during["recorded_days"]>calendar:raise ValueError("Impossible calendar coverage")
        no_row=calendar-during["recorded_days"]
        result.append({
            "site":SITE_NAMES[site],"year":year,
            "planting_date_published":start.isoformat(),
            "final_sample_date_NOT_commercial_harvest":end.isoformat(),
            "conservative_calendar_days_including_both_anchors":calendar,
            "registered_calendar_days_within_window":during["recorded_days"],
            "calendar_days_without_register_row_UNKNOWN":no_row,
            "W1_unknown_total_blanks_plus_unrecorded_days":during["W1"]["blank_UNKNOWN"]+no_row,
            "W2_unknown_total_blanks_plus_unrecorded_days":during["W2"]["blank_UNKNOWN"]+no_row,
            "windows":summaries,
        })
    return {
        "status":"M1_33_WUR_IRRIGATION_TRIAL_WINDOW_UNKNOWN_COVERAGE_NOT_WOFOST",
        "original_sha256":EXPECTED_SHA256 if strict else "SYNTHETIC_FIXTURE",
        "source_doi":"10.7910/DVN/1LC6W7","source_license":"CC BY-NC-SA 4.0",
        "source_sha_checked":strict,
        "development_year_2019_holdout_year_2020_preserved":True,
        "site_year_aggregate_counts":result,
        "blanks_or_absent_calendar_rows_are_NOT_zero":True,
        "source_unit_W1_W2_independently_confirmed":False,
        "source_blank_semantics_confirmed":False,
        "source_dates_or_water_doses_exported":0,
        "crop_yields_or_plot_ids_exported":0,
        "actual_crop_emergence_harvest_soil_genotype_verified":False,
        "wofost_irrigation_agromanagement_admissible":False,
        "wofost_crop_yield_predictions_executed":False,
    }

def inspect_original(blob):
    if hashlib.sha256(blob).hexdigest()!=EXPECTED_SHA256:
        raise ValueError("Original WUR irrigation SHA256 mismatch")
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    try:
        sheets={s.title.strip().lower():s for s in book.worksheets}
        if set(sheets)!={"meta","data"}:raise ValueError("Original workbook sheet names changed")
        safe_source_signatures(sheets["meta"],sheets["data"])
        it=sheets["data"].iter_rows(min_row=1,values_only=True)
        if next(it,None) is None:raise ValueError("No original data")
        return assess_rows(it)
    finally:book.close()

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--irrigation",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    x=p.parse_args()
    result=inspect_original(x.irrigation.read_bytes())
    x.out.mkdir(parents=True,exist_ok=True)
    (x.out/"m1_33_irrigation_trial_window_UNKNOWN_counts.json").write_text(
        json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps({"status":result["status"],
      "site_years":[{
         "site":v["site"],"year":v["year"],
         "calendar":v["conservative_calendar_days_including_both_anchors"],
         "source_in_window":v["registered_calendar_days_within_window"],
         "no_register_row_UNKNOWN":v["calendar_days_without_register_row_UNKNOWN"],
         "W1_UNKNOWN":v["W1_unknown_total_blanks_plus_unrecorded_days"],
         "W2_UNKNOWN":v["W2_unknown_total_blanks_plus_unrecorded_days"],
         "window_rows":{name:bucket["recorded_days"] for name,bucket in v["windows"].items()}
      } for v in result["site_year_aggregate_counts"]],
      "WOFOST_admissible":False},sort_keys=True))

if __name__=="__main__":main()
