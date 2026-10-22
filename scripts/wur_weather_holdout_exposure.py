"""M1-25: missing weather cells intersecting observed sampling endpoint.

ten Den et al., Harvard Dataverse v3.0, DOI 10.7910/DVN/1LC6W7,
CC BY-NC-SA 4.0. Publish COUNTS ONLY. Calendar-to-last-sample is
NOT emergence-to-harvest. WOFOST is never admitted from this audit.
"""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import io
import json
import math
from pathlib import Path

SOURCE = {
    "Lelystad": {
        "sha256": "ad66a23468a969b5795a656fb6989499e790a3e18179b9e3158f203b12ab8636",
        "last_sample_2019": "2019-10-31", "last_sample_2020": "2020-11-03"},
    "Vredepeel": {
        "sha256": "0bb40a5aa91ae6bfd02a8e8b4f825e4d834207884c1f2f44caa12fd0e88a1fd6",
        "last_sample_2019": "2019-10-10", "last_sample_2020": "2020-10-19"},
}
YEARS = {2018: 365, 2019: 365, 2020: 366}
NA = {"", "NA", "NAN", "NULL", "N/A", "-"}
REQUIRED_PCSE = ("IRRAD","TMIN","TMAX","VAP","WIND","RAIN","E0","ES0","ET0")

def parse_day(raw):
    if isinstance(raw, datetime): return raw.date()
    if isinstance(raw, date): return raw
    if isinstance(raw, str) and len(raw.strip()) == 10:
        try:
            result=date.fromisoformat(raw.strip())
            if raw.strip()==result.isoformat():return result
        except ValueError:pass
    return None

def value_state(v):
    if v is None or isinstance(v,str) and v.strip().upper() in NA:
        return "missing"
    if isinstance(v,bool):return "invalid"
    if isinstance(v,(float,int)) and math.isfinite(v):
        return "numeric"
    return "invalid"

def audit_sheet(ws,site,*,expected_years=YEARS):
    if site not in SOURCE:raise ValueError("Unknown original site")
    source_rows=list(ws.iter_rows(min_row=1,values_only=True))
    if not source_rows:raise ValueError("Empty weather source")
    visits=[]
    started=False
    for row in source_rows:
        day=parse_day(row[0]) if row else None
        if day is None:
            if started:raise ValueError("Undated row inside original weather series")
            continue
        started=True
        visits.append((day,row))
    if not visits:raise ValueError("Source contains no dated weather rows")
    dates=[d for d,_ in visits]
    if len(set(dates))!=len(dates):raise ValueError("Repeated daily weather identity")
    if dates!=sorted(dates) or (dates[-1]-dates[0]).days+1!=len(dates):
        raise ValueError("Missing or unordered weather days")
    if dict(Counter(d.year for d in dates))!=dict(expected_years):
        raise ValueError("Unexpected original weather calendar")
    maxcols=max(len(r) for r in source_rows)
    if maxcols<2:raise ValueError("Source weather columns absent")
    affected=[]
    for c in range(1,maxcols):
        by_year={}
        for year in (2019,2020):
            endpoint=date.fromisoformat(SOURCE[site]["last_sample_"+str(year)])
            numeric=missing=invalid=within=after=0
            for d,row in visits:
                if d.year!=year:continue
                state=value_state(row[c] if c<len(row) else None)
                if state=="numeric":
                    numeric+=1
                    continue
                if state=="missing":missing+=1
                else:invalid+=1
                if d<=endpoint:within+=1
                else:after+=1
            if numeric+missing+invalid!=expected_years[year]:
                raise ValueError("Incorrect weather availability balance")
            by_year[str(year)]={
                "calendar_days":expected_years[year],
                "numeric":numeric,"missing_na":missing,"invalid":invalid,
                "non_numeric_from_jan1_through_last_sample":within,
                "non_numeric_after_last_sample":after,
                "last_sample_not_confirmed_harvest":True}
        if any(d["missing_na"] or d["invalid"] for d in by_year.values()):
            affected.append({
                "source_column_1_based":c+1,
                "pcse_variable_identity_confirmed":False,
                "measurement_unit_confirmed":False,
                "coverage":by_year,
            })
    return {
        "site":site,
        "source_sha256":SOURCE[site]["sha256"],
        "original_weather_days":len(dates),
        "measurement_column_count":maxcols-1,
        "columns_with_missing_or_invalid_cells":affected,
        "model_ready":False,
    }

def audit_verified_files(blobs):
    if set(blobs)!=set(SOURCE):raise ValueError("Both original WUR XLSX required")
    for name, info in SOURCE.items():
        if hashlib.sha256(blobs[name]).hexdigest()!=info["sha256"]:
            raise ValueError("Original XLSX SHA-256 changed")
    from openpyxl import load_workbook
    results=[]
    for name in SOURCE:
        wb=load_workbook(io.BytesIO(blobs[name]),read_only=True,data_only=True)
        try:
            if len(wb.worksheets)!=1:raise ValueError("Source worksheet count changed")
            results.append(audit_sheet(wb.worksheets[0],name))
        finally:wb.close()
    return {
        "status":"M1_25_MISSINGNESS_VS_SAMPLING_ENDPOINT_NOT_WOFOST_VALIDATION",
        "doi":"10.7910/DVN/1LC6W7",
        "dataset_version":"3.0","license":"CC BY-NC-SA 4.0",
        "window_begin_is_Jan1_not_emergence":True,
        "window_end_last_sample_not_harvest":True,
        "original_daily_weather_values_exported":0,
        "original_plot_data_exported":0,
        "PCSE_field_mapping_approved":False,
        "weather_input_units_independently_verified":False,
        "soil_management_genotype_matching_verified":False,
        "field_wofost_run_admissible":False,
        "PCSE_fields_pending_provenance":list(REQUIRED_PCSE),
        "sites":results,
        "limits":[
            "The January-to-last-sample window does not assert actual emergence or commercial harvest.",
            "Source column position cannot be equated to any unconfirmed variable.",
            "Missing observations are never imputed or interpolated.",
            "Even complete numeric coverage does not certify PCSE semantics/units.",
            "WOFOST accuracy and Fe/Zn improvement have not been established.",
        ],
    }

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--lelystad",type=Path,required=True)
    p.add_argument("--vredepeel",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    a=p.parse_args()
    report=audit_verified_files({
        "Lelystad":a.lelystad.read_bytes(),"Vredepeel":a.vredepeel.read_bytes()})
    a.out.mkdir(parents=True,exist_ok=True)
    (a.out/"m1_25_weather_holdout_exposure_counts.json").write_text(
        json.dumps(report,sort_keys=True,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":report["status"],"sites":[{
        "site":site["site"],"missing":[{"column":c["source_column_1_based"],
        "by_year":{y:{"through_last_sample":x["non_numeric_from_jan1_through_last_sample"],
                      "after_last_sample":x["non_numeric_after_last_sample"]}
                   for y,x in c["coverage"].items()}}
        for c in site["columns_with_missing_or_invalid_cells"]]}
        for site in report["sites"]],"model_admitted":False},sort_keys=True))

if __name__=="__main__":main()
