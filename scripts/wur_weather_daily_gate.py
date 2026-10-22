"""M1-22: original Dutch weather chronology + strict PCSE weather unit contract.

Provenance: ten Den et al. (2024), DOI 10.7910/DVN/1LC6W7,
Harvard Dataverse v3.0, CC BY-NC-SA 4.0.
No original workbook, cell, weather series, worksheet label or station
coordinates are published. No WOFOST run or model validation is authorized.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from scripts.wur_weather_readiness import FILES, LICENSE

METRIC_FIELDS = ("IRRAD", "TMIN", "TMAX", "VAP", "RAIN", "WIND", "E0", "ES0", "ET0")
FIELDS = ("DAY",) + METRIC_FIELDS
REQUIRED_UNITS = {
    "IRRAD": {"J/m2/day": 1.0, "kJ/m2/day": 1_000.0, "MJ/m2/day": 1_000_000.0},
    "TMIN": {"degC": 1.0}, "TMAX": {"degC": 1.0},
    "VAP": {"hPa": 1.0, "kPa": 10.0},
    "RAIN": {"cm/day": 1.0, "mm/day": 0.1},
    "WIND": {"m/s_at_2m": 1.0},
    "E0": {"cm/day": 1.0, "mm/day": 0.1},
    "ES0": {"cm/day": 1.0, "mm/day": 0.1},
    "ET0": {"cm/day": 1.0, "mm/day": 0.1},
}
RANGES = {"IRRAD": (0, 40e6), "TMIN": (-50, 60), "TMAX": (-50, 60),
          "VAP": (0.06, 199.3), "RAIN": (0, 25), "WIND": (0, 50),
          "E0": (0, 2.5), "ES0": (0, 2.5), "ET0": (0, 2.5)}


def parse_day(value):
    """No inference from ambiguous date, locale or untyped Excel serial."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _date_column(sheet):
    """Inspect only rows after the published top-nine-row metadata section."""
    counts = Counter()
    visits = {}
    nonempty = Counter()
    for row in sheet.iter_rows(min_row=10, values_only=True):
        for i, value in enumerate(row):
            if value is not None:
                nonempty[i] += 1
            d = parse_day(value)
            if d is not None and d.year in (2019, 2020):
                counts[i] += 1
                visits.setdefault(i, []).append(d)
    # Date axis must be unique, nontrivial and cover most non-empty cells.
    if not counts:
        return {"date_column_status": "NONE_UNAMBIGUOUS", "matching_date_column_count": 0,
                "parsed_dates_2019": 0, "parsed_dates_2020": 0,
                "distinct_dates": 0, "duplicate_date_rows": 0,
                "missing_calendar_days_within_observed_window": None,
                "date_axis_ready_for_model": False}
    best = max(counts.values())
    possible = [i for i, n in counts.items()
                if n == best and n >= 20 and (n / nonempty[i]) >= 0.90]
    if len(possible) != 1:
        return {"date_column_status": "AMBIGUOUS_OR_TOO_SPARSE",
                "matching_date_column_count": len(possible),
                "parsed_dates_2019": 0, "parsed_dates_2020": 0,
                "distinct_dates": 0, "duplicate_date_rows": 0,
                "missing_calendar_days_within_observed_window": None,
                "date_axis_ready_for_model": False}
    series = visits[possible[0]]
    by_year = Counter(d.year for d in series)
    unique = set(series)
    calendar_gaps = 0
    for year in (2019, 2020):
        subset = sorted(d for d in unique if d.year == year)
        if subset:
            calendar_gaps += (subset[-1] - subset[0]).days + 1 - len(subset)
    dup = len(series) - len(unique)
    return {"date_column_status": "UNIQUE_DATE_LIKE_COLUMN_UNITS_UNCONFIRMED",
            "matching_date_column_count": 1,
            "parsed_dates_2019": by_year[2019],
            "parsed_dates_2020": by_year[2020],
            "distinct_dates": len(unique),
            "duplicate_date_rows": dup,
            "missing_calendar_days_within_observed_window": calendar_gaps,
            "date_axis_ready_for_model": False}


def inspect_original_weather(blobs):
    """Use original checksums; public result has no cell values or original names."""
    from openpyxl import load_workbook
    if set(blobs) != set(FILES):
        raise ValueError("Both official weather files are mandatory")
    sites = []
    for site in FILES:
        if hashlib.sha256(blobs[site]).hexdigest() != FILES[site]["sha256"]:
            raise ValueError("Original weather XLSX checksum mismatch")
        wb = load_workbook(io.BytesIO(blobs[site]), read_only=True, data_only=True)
        try:
            if not wb.worksheets:
                raise ValueError("No source Excel sheets")
            sheets = []
            for ws in wb.worksheets:
                sheets.append({"sheet_position": len(sheets) + 1,
                               "reported_rows": ws.max_row, "reported_columns": ws.max_column,
                               **_date_column(ws)})
            sites.append({"site": site, "file_id": FILES[site]["id"],
                          "source_sha256": FILES[site]["sha256"], "sheets": sheets})
        finally:
            wb.close()
    return {
        "status": "ORIGINAL_WUR_WEATHER_CHRONOLOGY_AUDITED_NOT_WOFOST_READY",
        "doi": "10.7910/DVN/1LC6W7",
        "dataset_version": "3.0", "license": LICENSE,
        "source_worksheet_rows_exported": 0, "source_daily_weather_values_exported": 0,
        "original_excel_files_redistributed": False,
        "date_semantics_and_frequency_confirmed_by_source_dictionary": False,
        "source_variables_units_and_conversion_provenance_approved": False,
        "actual_field_soil_irrigation_and_cultivar_model_admitted": False,
        "field_wofost_run_admissible": False,
        "sites": sites,
        "limits": [
            "Top nine rows are treated as workbook metadata following ten Den et al. (2024); this is a structural convention, not a verified row-level mapping.",
            "Excel serial values without date type and non-ISO date strings are not guessed.",
            "A single apparent date column does NOT establish full source coverage or unit semantics.",
            "Relative humidity percent is NOT vapor pressure (hPa).",
            "Makkink reference evapotranspiration cannot be silently rebranded Penman E0/ES0.",
            "Weather series are not redistributed; original CC BY-NC-SA constraints remain.",
            "No WOFOST predictive accuracy, Fe/Zn, nutrition, safety or genetic effects inferred.",
        ],
    }


def normalize_for_pcse(rows, *, units, unit_provenance_approved=False):
    """Pure daily converter for LOCAL/synthetic values, NOT scientific admission.

    Caller must supply every weather variable and per-field source unit.
    Source-specific calibration/validation is a *separate* approval gate.
    """
    if set(units) != set(METRIC_FIELDS):
        raise ValueError("Missing or extra units in daily weather mapping")
    factors = {}
    for key in METRIC_FIELDS:
        if units[key] not in REQUIRED_UNITS[key]:
            raise ValueError(f"Unsupported/ambiguous meteorological unit for {key}")
        factors[key] = REQUIRED_UNITS[key][units[key]]
    if not rows:
        raise ValueError("Empty meteorological time series")
    daily = []
    days = set()
    for row in rows:
        if set(row) != set(FIELDS):
            raise ValueError("Missing or extra daily weather columns")
        day = parse_day(row["DAY"])
        if day is None:
            raise ValueError("Ambiguous source day: require ISO YYYY-MM-DD or typed date")
        if day in days:
            raise ValueError("Duplicate meteorological day is not independent data")
        days.add(day)
        record = {"DAY": day}
        for key in METRIC_FIELDS:
            value = row[key]
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"Missing or nonnumeric daily weather {key}")
            x = float(value) * factors[key]
            low, high = RANGES[key]
            if not math.isfinite(x) or not low <= x <= high:
                raise ValueError(f"Nonfinite/out-of-range daily weather {key}")
            record[key] = x
        if record["TMIN"] > record["TMAX"]:
            raise ValueError("Minimum temperature exceeds maximum")
        daily.append(record)
    daily.sort(key=lambda row: row["DAY"])
    if (daily[-1]["DAY"] - daily[0]["DAY"]).days + 1 != len(daily):
        raise ValueError("Missing daily weather dates; implicit interpolation forbidden")
    # Conversion is a necessary technical contract, NOT approval of source meaning.
    return {"status": "UNIT_NORMALIZATION_ONLY_NOT_MODEL_ADMISSION",
            "unit_provenance_approved_by_independent_review": bool(unit_provenance_approved),
            "source_model_climate_alignment_verified": False,
            "field_wofost_run_admissible": False,
            "pcse_daily_records_IN_MEMORY_DO_NOT_EXPORT_OR_PUBLISH": daily}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--lelystad", type=Path, required=True)
    parser.add_argument("--vredepeel", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    r = inspect_original_weather({"Lelystad": args.lelystad.read_bytes(),
                                  "Vredepeel": args.vredepeel.read_bytes()})
    args.out.mkdir(parents=True, exist_ok=True)
    (args.out / "m1_22_weather_chronology_metadata_only.json").write_text(
        json.dumps(r, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": r["status"], "source_values_exported": 0,
                      "weather_model_ready": r["field_wofost_run_admissible"],
                      "site_sheet_summaries": [{"site":x["site"],"sheets":x["sheets"]}
                                               for x in r["sites"]]}, sort_keys=True))


if __name__ == "__main__":
    main()
