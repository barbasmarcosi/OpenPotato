"""M1-21: audit original WUR Dutch weather XLSX without deriving crop weather.

Original workbooks CC BY-NC-SA 4.0 (ten Den et al. 2024;
10.7910/DVN/1LC6W7). Output structural metadata only: no cell values,
sheet names, station coordinates, climatic measurements, or time series.
Never claim these XLSX are validated PCSE/WOFOST daily drivers.
"""
from __future__ import annotations
import argparse
import hashlib
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook

FILES = {
    "Lelystad": {"id": 10284792,
                 "sha256": "ad66a23468a969b5795a656fb6989499e790a3e18179b9e3158f203b12ab8636"},
    "Vredepeel": {"id": 10284786,
                  "sha256": "0bb40a5aa91ae6bfd02a8e8b4f825e4d834207884c1f2f44caa12fd0e88a1fd6"},
}
DOI = "10.7910/DVN/1LC6W7"
LICENSE = "CC BY-NC-SA 4.0"
# Header matching is suggestive, never confirmation of units or meaning.
TOKENS = {
    "temporal": (r"\bdate\b", r"\bdatum\b", r"\byear\b", r"\bday\b", r"\bdoy\b"),
    "radiation": (r"radiat", r"\bglob", r"\birrad", r"\bsrad"),
    "precipitation": (r"\brain", r"precip", r"neerslag"),
    "temperature": (r"\btemp", r"temperature", r"tmin", r"tmax"),
    "humidity": (r"humid", r"damp", r"vapour", r"vapor"),
    "wind": (r"\bwind",),
}


def describe_workbook(blob: bytes) -> dict:
    """Inspect IN-MEMORY XLSX without exporting cell values."""
    try:
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    except Exception as ex:
        raise ValueError("Invalid or unreadable XLSX source") from ex
    summary = []
    try:
        if not wb.worksheets:
            raise ValueError("Workbook has no worksheets")
        for ws in wb.worksheets:
            header_rows = []
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, 12), values_only=True):
                header_rows.append(row)
            labels = [str(value).strip().lower() for row in header_rows
                      for value in row if isinstance(value, str)]
            matches = {key: sum(any(re.search(pattern, label) for pattern in patterns)
                                for label in labels) for key, patterns in TOKENS.items()}
            summary.append({
                "sheet_position_1_based": len(summary) + 1,
                "reported_max_rows": ws.max_row,
                "reported_max_columns": ws.max_column,
                "header_candidate_rows_inspected": len(header_rows),
                "candidate_header_token_counts_unconfirmed_units": matches,
            })
    finally:
        wb.close()
    return {"worksheet_count": len(summary), "structure_only": summary}


def audit_sources(blobs: dict[str, bytes]) -> dict:
    if set(blobs) != set(FILES):
        raise ValueError("Both source weatherbooks required, no alternative stations")
    sites = []
    for name in FILES:
        sha = hashlib.sha256(blobs[name]).hexdigest()
        if sha != FILES[name]["sha256"]:
            raise ValueError("Source weather XLSX SHA256 changed; manual rebaseline required")
        sites.append({"site": name, "dataverse_file_id": FILES[name]["id"],
                      "source_sha256": sha, **describe_workbook(blobs[name])})
    return {
        "status": "VERIFIED_WEATHER_WORKBOOK_STRUCTURE_ONLY_NOT_MODEL_READY",
        "source_doi": DOI, "source_version": "3.0", "source_license": LICENSE,
        "source_license_uri": "https://creativecommons.org/licenses/by-nc-sa/4.0/",
        "sites": sites,
        "original_weather_cell_values_exported": 0,
        "raw_daily_weather_records_exported": 0,
        "original_workbooks_redistributed": False,
        "model_driver_units_and_daily_frequency_confirmed": False,
        "crop_parameter_soil_irrigation_join_confirmed": False,
        "ready_to_run_field_wofost": False,
        "limits": [
            "Header labels matched approximately and do not validate a column's units, source or frequency.",
            "Excel formulas read with data_only may lack caches; numeric data were NOT derived.",
            "No yield, tuber nutrients or Fe/Zn prediction; no WOFOST validation or calibration.",
            "Requires expert review of original variable dictionary, time period, unit conversions and site matching.",
            "The original source has CC BY-NC-SA 4.0 license; preserve NC and SA conditions.",
        ],
    }


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--lelystad", type=Path, required=True)
    p.add_argument("--vredepeel", type=Path, required=True)
    p.add_argument("--out", type=Path, required=True)
    a = p.parse_args()
    report = audit_sources({"Lelystad": a.lelystad.read_bytes(),
                            "Vredepeel": a.vredepeel.read_bytes()})
    a.out.mkdir(parents=True, exist_ok=True)
    (a.out / "wur_weather_workbook_structure.json").write_text(
        json.dumps(report, sort_keys=True, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "weatherbook_count": len(report["sites"]),
                      "ready_to_run_field_wofost": False}, sort_keys=True))


if __name__ == "__main__":
    main()
