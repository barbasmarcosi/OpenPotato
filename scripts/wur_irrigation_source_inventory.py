"""M1-29: WUR ORIGINAL irrigation workbook metadata-only audit.

Do not treat a row or numeric cell as a plot-level irrigation event without
verified columns, unit and physical plot linkage. Licensed CC BY-NC-SA 4.0.
Only counts/controlled boolean hints leave the ephemeral runner.
"""
from __future__ import annotations
import argparse
from collections import Counter
from datetime import date,datetime
import hashlib
import io
import json
import re
from pathlib import Path

SOURCE_DOI="10.7910/DVN/1LC6W7"
SOURCE_ID=10284791
SOURCE_NAME="Irrigation.xlsx"
LICENSE="CC BY-NC-SA 4.0"
# Digest to be independently pinned after initial source provenance review.
EXPECTED_SHA256=None
TAGS={
    "date":r"\bdate\b|\bdatum\b|\bday\b",
    "year":r"\byear\b|\bjaar\b",
    "location":r"\blocation\b|\bsite\b|\blocatie\b",
    "physical_plot":r"\bplot(?:number|id)?\b|\bfield\s*id\b",
    "irrigation_volume_or_amount":r"\bmm\b|\bamount\b|\bvolume\b|\birrigation\b|\bwater\s*applied\b",
    "treatment":r"\btreatment\b|\btrt\b",
}
def parse_date(x):
    if isinstance(x,datetime):return x.date()
    if isinstance(x,date):return x
    if isinstance(x,str):
        try:
            y=date.fromisoformat(x.strip())
            if x.strip()==y.isoformat():return y
        except ValueError:pass
    return None

def inspect_sheet(ws):
    # Real source cells stay transient; output only source structure.
    rows=0
    finite_numbers=0
    year_dates=Counter()
    known_tags=Counter()
    samples=0
    for row in ws.iter_rows(values_only=True):
        rows+=1
        if rows>200000:raise ValueError("Source sheet row safety limit exceeded")
        for value in row:
            if type(value) in (int,float) and not isinstance(value,bool):
                finite_numbers+=1
            d=parse_date(value)
            if d and d.year in (2019,2020):year_dates[str(d.year)]+=1
            if rows<=30 and isinstance(value,str):
                samples+=1
                for tag,pattern in TAGS.items():
                    if re.search(pattern,value.strip(),re.IGNORECASE):
                        known_tags[tag]+=1
    return {
        "physical_rows_inspected":rows,
        "declared_max_columns":ws.max_column,
        "numeric_cells_not_a_validated_water_amount":finite_numbers,
        "typed_or_ISO_2019_2020_dates_found_by_year":dict(year_dates),
        "controlled_source_vocabulary_hints":dict(known_tags),
        "original_header_text_exported":False,
        "original_calendar_dates_exported":False,
        "water_doses_or_event_rows_exported":0,
        "source_schema_or_units_approved":False,
    }

def inspect_irrigation(blob:bytes,*,expected_sha=None):
    if not blob:raise ValueError("Empty original irrigation source")
    sha=hashlib.sha256(blob).hexdigest()
    pinned=expected_sha if expected_sha is not None else EXPECTED_SHA256
    if pinned is not None and sha!=pinned:
        raise ValueError("Pinned original irrigation SHA256 mismatch")
    from openpyxl import load_workbook
    try:workbook=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    except Exception as ex:raise ValueError("Original irrigation XLSX invalid") from ex
    try:
        sheets=[]
        for idx,ws in enumerate(workbook.worksheets):
            sheet_kind=ws.title.strip().lower()
            if sheet_kind not in ("meta","data"):sheet_kind="other"
            sheets.append({"sheet_position_one_based":idx+1,
                           "whitelisted_type_if_recognized":sheet_kind,
                           **inspect_sheet(ws)})
    finally:workbook.close()
    if not sheets:raise ValueError("No original worksheet available")
    return {
        "status":"M1_29_ORIGINAL_IRRIGATION_WORKBOOK_INVENTORIED_NO_EVENT_LINKAGE",
        "source_doi":SOURCE_DOI,"source_file_id":SOURCE_ID,
        "source_filename":SOURCE_NAME,"source_sha256":sha,
        "checksum_independently_pinned_before_publication":pinned is not None,
        "source_license":LICENSE,
        "sheet_count":len(sheets),"sheet_structures":sheets,
        "original_sheet_names_or_cell_values_exported":False,
        "plot_level_irrigation_dates_or_amounts_verified":False,
        "irrigation_source_unit_mm_verified":False,
        "plot_or_treatment_to_event_join_verified":False,
        "real_wofost_irrigation_schedule_ready":False,
        "warnings":[
            "Source workbook may include treatment metadata and dated watering but cell counts cannot prove an event.",
            "Date and mm header clues are NOT validated column mappings.",
            "No source rows, dates, irrigation doses, plot numbers or individual records are redistributed.",
            "Do not substitute irrigation treatment labels for actual mm and date events.",
        ],
    }

def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--source",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    p.add_argument("--expected-sha256",default=None)
    a=p.parse_args()
    result=inspect_irrigation(a.source.read_bytes(),expected_sha=a.expected_sha256)
    a.out.mkdir(parents=True,exist_ok=True)
    (a.out/"m1_29_irrigation_workbook_structure_only.json").write_text(
        json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps({"status":result["status"],"digest":result["source_sha256"],
                      "sheets":[{"kind":x["whitelisted_type_if_recognized"],
                                 "rows":x["physical_rows_inspected"],
                                 "columns":x["declared_max_columns"],
                                 "year_date_counts":x["typed_or_ISO_2019_2020_dates_found_by_year"],
                                 "controlled_hints":x["controlled_source_vocabulary_hints"]}
                                for x in result["sheet_structures"]],
                      "real_irrigation_event_join_ready":False},sort_keys=True))
if __name__=="__main__":main()
